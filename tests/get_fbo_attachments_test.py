import unittest
import responses
import requests_mock
import requests
from fpdf import FPDF
from docx import Document
from bs4 import BeautifulSoup
import sys
import os
sys.path.append( os.path.dirname( os.path.dirname( os.path.abspath(__file__) ) ) )
from utils.get_fbo_attachments import FboAttachments
from fixtures import nightly_data, fedconnect

class FboAttachmentsTestCase(unittest.TestCase):

    def setUp(self):
        self.maxDiff = None
        self.fake_fbo_url = 'https://www.fbo.gov/fake'
        self.fboa = FboAttachments(nightly_data = nightly_data.nightly_data)
        text = "This is a test"
        cwd = os.getcwd()
        if 'fbo-scraper' in cwd:
            i = cwd.find('fbo-scraper')
            root_path = cwd[:i+len('fbo-scraper')]
        else:
            i = cwd.find('root')
            root_path = cwd
        temp_outfile_path = os.path.join(root_path, 'temp_test_file')
        with open(temp_outfile_path, 'w') as f:
            f.write(text)
        self.temp_outfile_path = temp_outfile_path

        txt_text = "This is a test"
        temp_outfile_path_txt = os.path.join(root_path, 'temp_test_file_txt.txt')
        with open(temp_outfile_path_txt, 'w') as f:
            f.write(txt_text)
        self.temp_outfile_path_txt = temp_outfile_path_txt

        pdf = FPDF()
        pdf.add_page()
        pdf.set_xy(0, 0)
        pdf.set_font('arial', 'B', 13.0)
        pdf.cell(ln=0, h=5.0, align='L', w=0, txt="This is a test", border=0)
        temp_outfile_path_pdf  = os.path.join(root_path, 'test.pdf')
        pdf.output(temp_outfile_path_pdf, 'F')
        self.temp_outfile_path_pdf = temp_outfile_path_pdf

        document_docx = Document()
        document_docx.add_heading("This is a test", 0)
        document_docx.save('test.docx')
        temp_outfile_path_docx = os.path.join(root_path, 'test.docx')
        temp_outfile_path_doc = os.path.join(root_path, 'tests/fixtures/test.doc')
        self.temp_outfile_path_docx = temp_outfile_path_docx
        self.temp_outfile_path_doc = temp_outfile_path_doc


    def tearDown(self):
        self.fboa = None
        self.fake_fbo_url = None
        if os.path.exists(self.temp_outfile_path):
            os.remove(self.temp_outfile_path)
        if os.path.exists(self.temp_outfile_path_txt):
            os.remove(self.temp_outfile_path_txt)
        if os.path.exists(self.temp_outfile_path_pdf):
            os.remove(self.temp_outfile_path_pdf)
        if os.path.exists(self.temp_outfile_path_docx):
            os.remove(self.temp_outfile_path_docx)


    @requests_mock.Mocker()
    def test_get_divs(self, mock_request):
        body_with_div = b'''
                            <div class="notice_attachment_ro notice_attachment_last">
                            <div>
                        <div class="file"><input type="hidden" name="dnf_class_values[procurement_notice_archive][packages][5][files][0][file][0][preview]" value="&lt;a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'&gt;FA852618Q0033_______0001.pdf&lt;/a&gt;"><a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'>FA852618Q0033_______0001.pdf</a> (16.54 Kb)</div>
                        </div>
                            <div><span class="label">Description:</span> Amendment 0001</div>	</div>


                            </div><!-- widget -->

                            </div>
        '''
        mock_request.register_uri('GET',
                                  url=self.fake_fbo_url,
                                  content=body_with_div,
                                  status_code = 200)
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = ['expected div in here']
        self.assertEqual(len(result), len(expected))

    @requests_mock.Mocker()
    def test_get_divs_wrong_url(self, mock_request):
        mock_request.register_uri('GET',
                                  url=self.fake_fbo_url,
                                  text='No divs in here',
                                  status_code = 200)
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))

    @requests_mock.Mocker()
    def test_get_divs_non200_url(self, mock_request):
        mock_request.register_uri('GET',
                                  url=self.fake_fbo_url,
                                  status_code=404)
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))

    @requests_mock.Mocker()
    def test_get_divs_connection_error(self, mock_request):
        mock_request.register_uri('GET',
                                  url=self.fake_fbo_url,
                                  exc=requests.exceptions.ConnectionError)
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))

    @requests_mock.Mocker()
    def test_get_neco_navy_mil_attachment_urls_singleton(self, mock_request):
        body = b'''
                <table id="tbl7" border="0" style="width:600px;">
                    <tbody><tr id="dwnld2_row">
                        <td class="tbl_hdr" align="right" style="width:150px;">Download File: </td><td class="tbl_itm_sm" align="left">&nbsp;<a id="dwnld2_hypr" href="/upload/N00189/N0018919Q0353Combined_Synopsis_Solicitation.docx" target="_blank">N00189/N0018919Q0353Combined_Synopsis_Solicitation.docx</a></td>
                    </tr>
                </tbody></table>
                '''
        mock_request.register_uri('GET',
                                  url = self.fake_fbo_url,
                                  status_code = 200,
                                  content = body)
        result = self.fboa.get_neco_navy_mil_attachment_urls(self.fake_fbo_url)
        expected = ["https://www.neco.navy.mil/upload/N00189/N0018919Q0353Combined_Synopsis_Solicitation.docx"]
        self.assertListEqual(result, expected)

    @requests_mock.Mocker()
    def test_get_neco_navy_mil_attachment_urls_multiple(self, mock_request):
        body = b'''
                <tbody><tr id="dwnld2_row">
                        <td class="tbl_hdr" align="right" style="width:150px;">Download File: </td><td class="tbl_itm_sm" align="left">&nbsp;<a id="dwnld2_hypr" href="/upload/N00406/N0040619Q0062N00406-19-Q-0062_SOLICITATION.pdf" target="_blank">N00406/N0040619Q0062N00406-19-Q-0062_SOLICITATION.pdf</a></td>
                    </tr><tr id="dwnld3_row">
                        <td class="tbl_hdr" align="right" style="width:150px;">Download File: </td><td class="tbl_itm_sm" align="left">&nbsp;<a id="dwnld3_hypr" href="/upload/N00406/N0040619Q0062252.211-7003_ITEM_UNIQUE_IDENTIFICATION__VALUATION.docx" target="_blank">N00406/N0040619Q0062252.211-7003_ITEM_UNIQUE_IDENTIFICATION__VALUATION.docx</a></td>
                    </tr><tr id="dwnld4_row">
                        <td class="tbl_hdr" align="right" style="width:150px;">Download File: </td><td class="tbl_itm_sm" align="left">&nbsp;<a id="dwnld4_hypr" href="/upload/N00406/N0040619Q0062NAVSUP_FACTS-SP_Shipping_information.docx" target="_blank">N00406/N0040619Q0062NAVSUP_FACTS-SP_Shipping_information.docx</a></td>
                    </tr>
                </tbody>
                '''
        mock_request.register_uri('GET',
                                  url = self.fake_fbo_url,
                                  status_code = 200,
                                  content = body)
        result = self.fboa.get_neco_navy_mil_attachment_urls(self.fake_fbo_url)
        expected = ['https://www.neco.navy.mil/upload/N00406/N0040619Q0062N00406-19-Q-0062_SOLICITATION.pdf',
                    'https://www.neco.navy.mil/upload/N00406/N0040619Q0062252.211-7003_ITEM_UNIQUE_IDENTIFICATION__VALUATION.docx',
                    'https://www.neco.navy.mil/upload/N00406/N0040619Q0062NAVSUP_FACTS-SP_Shipping_information.docx']
        self.assertListEqual(result, expected)

    @requests_mock.Mocker()
    def test_get_neco_navy_mil_attachment_urls_connection_error(self, mock_request):
        mock_request.register_uri('GET',
                                  url = self.fake_fbo_url,
                                  exc = requests.exceptions.ConnectionError)
        result = self.fboa.get_neco_navy_mil_attachment_urls(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))

    @requests_mock.Mocker()
    def test_get_neco_navy_mil_attachment_urls_non200_url(self, mock_request):
        mock_request.register_uri('GET',
                                  url = self.fake_fbo_url,
                                  status_code = 404)
        result = self.fboa.get_neco_navy_mil_attachment_urls(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))

    def test_insert_attachments(self):
        text = "This is a test"
        notice = {'a': '1', 'b': '2'}
        with open(self.temp_outfile_path, 'w') as f:
            f.write(text)
        file_list = [(self.temp_outfile_path,
                      self.fake_fbo_url)]
        result = self.fboa.insert_attachments(file_list, notice)
        expected = {'a': '1',
                    'b': '2',
                    'attachments':[{'machine_readable': True,
                                    'text':text,
                                    'url':self.fake_fbo_url,
                                    'prediction':None,
                                    'decision_boundary':None,
                                    'validation':None,
                                    'trained':False}]
        }
        self.assertDictEqual(result, expected)

    @requests_mock.Mocker()
    def test_size_check(self, mock_request):
        mock_request.register_uri('HEAD',
                                  self.fake_fbo_url, 
                                  headers = {'Content-Length': '800'}, 
                                  status_code = 200)
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = True
        self.assertEqual(result, expected)

    @requests_mock.Mocker()
    def test_size_check_non200_url(self, mock_request):
        mock_request.register_uri('HEAD',
                                  url = self.fake_fbo_url,
                                  headers = {'Content-Length': '800'}, 
                                  status_code=404)
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = False
        self.assertEqual(result, expected)

    @requests_mock.Mocker()
    def test_size_check_connection_error(self, mock_request):
        mock_request.register_uri('HEAD',
                                  url = self.fake_fbo_url,
                                  exc = requests.exceptions.ConnectionError)
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = False
        self.assertEqual(result, expected)

    @requests_mock.Mocker()
    def test_size_check_redirect_true(self, mock_request):
        body = 'This is less than 500MB'
        redirect_location = 'https://www.fbo.gov/fakeredirect'
        mock_request.register_uri('HEAD',
                                  url = self.fake_fbo_url,
                                  status_code = 302,
                                  text = body,
                                  headers = {'Content-Type': 'application/pdf', 
                                                    'Content-Length': str(len(body)),
                                                    'Location': redirect_location})
        mock_request.register_uri('HEAD',
                                  url = redirect_location,
                                  status_code = 200,
                                  text = body,
                                  headers = {'Content-Type': 'application/pdf', 
                                             'Content-Length': str(len(body))})
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = True
        self.assertEqual(result, expected)

    @requests_mock.Mocker()
    def test_size_check_redirect_false(self, mock_request):
        body = 'body'
        big_body = "*"*600000000
        redirect_location = 'https://www.fbo.gov/fakeredirect'
        mock_request.register_uri('HEAD',
                                  url = self.fake_fbo_url,
                                  status_code = 302,
                                  text = body,
                                  headers = {'Content-Type': 'application/pdf', 
                                             'Content-Length': str(len(body)),
                                             'Location': redirect_location})
        mock_request.register_uri('HEAD',
                                  url = redirect_location,
                                  status_code = 200,
                                  text = big_body,
                                  headers = {'Content-Type': 'application/pdf', 
                                             'Content-Length': str(len(big_body))})
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = False
        self.assertEqual(result, expected)

    def test_get_filename_from_cd(self):
        content_disposition = 'attachment; filename="test.doc"'
        result = self.fboa.get_filename_from_cd(content_disposition)
        expected = 'test.doc'
        self.assertEqual(result, expected)

    def test_get_filename_from_bad_cd(self):
        content_disposition = 'no file name in here!'
        result = self.fboa.get_filename_from_cd(content_disposition)
        expected = None
        self.assertEqual(result, expected)

    def test_get_file_name(self):
        attachment_url = 'https://test.pdf'
        content_type = 'application/pdf'
        result = self.fboa.get_file_name(attachment_url, content_type)
        expected = 'test.pdf'
        self.assertEqual(result, expected)

    def test_get_attachment_text_txt(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_txt, 'url')
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_attachment_text_pdf(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_pdf, 'url')
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_attachment_text_docx(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_docx, 'url')
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_attachment_text_doc(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_doc, 'url')
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_attachment_url_from_div(self):
        div = '<a href="/utils/view?id=798e26de983ca76f9075de687047445a"\
               target="_blank" title="Download/View FD2060-17-33119_FORM_158_00.pdf"\
               class="file">FD2060-17-33119_FORM_158_00.pdf</a>'
        div = BeautifulSoup(div, "html.parser")
        result, is_neco_navy_mil = self.fboa.get_attachment_url_from_div(div)
        expected = ['https://www.fbo.gov/utils/view?id=798e26de983ca76f9075de687047445a']
        with self.subTest():
            self.assertListEqual(result, expected)
        with self.subTest():
            self.assertFalse(is_neco_navy_mil)

    def test_get_attachment_url_from_div_space(self):
        '''
        Test parsing of the oddly formatted divs
        '''
        div = '<a href="http://  https://www.thisisalinktoanattachment.docx"\
               target="_blank" title="Download/View FD2060-17-33119_FORM_158_00.pdf"\
               class="file">FD2060-17-33119_FORM_158_00.pdf</a>'
        div = BeautifulSoup(div, "html.parser")
        result, is_neco_navy_mil = self.fboa.get_attachment_url_from_div(div)
        expected = ['https://www.thisisalinktoanattachment.docx']
        with self.subTest():
            self.assertListEqual(result, expected)
        with self.subTest():
            self.assertFalse(is_neco_navy_mil)

    def test_write_attachments(self):
        #TODO don't rely on the url in the div below continuing to exist
        body_with_div = b'''
                            <div class="notice_attachment_ro notice_attachment_last">
                            <div>
                        <div class="file"><input type="hidden" name="dnf_class_values[procurement_notice_archive][packages][5][files][0][file][0][preview]" value="&lt;a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'&gt;FA852618Q0033_______0001.pdf&lt;/a&gt;"><a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'>FA852618Q0033_______0001.pdf</a> (16.54 Kb)</div>
                        </div>
                            <div><span class="label">Description:</span> Amendment 0001</div>	</div>
                            </div><!-- widget -->
                            </div>
                         '''
        soup = BeautifulSoup(body_with_div, "html.parser")
        attachment_divs = soup.find_all('div', {"class": "notice_attachment_ro"})
        result = self.fboa.write_attachments(attachment_divs)
        attachment_path = os.path.join(os.getcwd(),'attachments','FA852618Q0033_______0001.pdf')
        expected = [(attachment_path, 
                     'https://www.fbo.gov/utils/view?id=d95550fb782f53357ed65db571ef9186')]
        for file_url_tup in result:
            file, _ = file_url_tup
            os.remove(file)
        self.assertEqual(result, expected)

    def test_get_post_payload(self):
        a_tag = fedconnect.get_a_tag()
        soup = fedconnect.get_fedconnect_soup()
        result = self.fboa.get_post_payload(a_tag, soup)
        expected = fedconnect.expected_payload
        self.assertEqual(result, expected)

    
if __name__ == '__main__':
    unittest.main()