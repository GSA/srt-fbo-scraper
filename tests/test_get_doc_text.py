import os
import shutil
import sys

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import unittest
from unittest.mock import patch
from zipfile import ZipFile, BadZipfile

from fpdf import FPDF
from docx import Document
import textract

from fbo_scraper.get_doc_text import get_doc_text


class GetDocTextTestCase(unittest.TestCase):
    def setUp(self):
        self.maxDiff = None

        # set up output dir
        out_path = "test_attachments"
        if not os.path.exists(out_path):
            os.mkdir(out_path)
        self.abs_out_path = os.path.abspath(out_path)

        # refer to fixture
        if "fixtures" in os.getcwd():
            self.temp_outfile_path_doc = os.path.join("test.doc")
        elif os.path.exists(
            os.path.join(os.getcwd(), "tests", "fixtures")
        ):  # running in pycharm
            self.temp_outfile_path_doc = os.path.join(
                os.getcwd(), "tests", "fixtures", "test.doc"
            )
        else:
            self.temp_outfile_path_doc = os.path.join("fixtures", "test.doc")

        if os.getenv("CI"):
            self.temp_outfile_path_doc = "/root/project/tests/fixtures/test.doc"

        # create a mock file w/o extension
        text = "This is a test"
        self.temp_outfile_path = os.path.join(self.abs_out_path, "temp_test_file")
        with open(self.temp_outfile_path, "w") as f:
            f.write(text)

        # create mock txt file
        self.temp_outfile_path_txt = os.path.join(
            self.abs_out_path, "temp_test_file_txt.txt"
        )
        with open(self.temp_outfile_path_txt, "w") as f:
            f.write(text)
        

        # create mock pdf
        pdf = FPDF()
        pdf.add_page()
        pdf.set_xy(0, 0)
        pdf.set_font("arial", "B", 13.0)
        pdf.cell(ln=0, h=5.0, align="L", w=0, txt=text, border=0)
        self.temp_outfile_path_pdf = os.path.join(self.abs_out_path, "test.pdf")
        pdf.output(self.temp_outfile_path_pdf, "F")

        # create a docx
        document_docx = Document()
        document_docx.add_paragraph(text)
        self.temp_outfile_path_docx = os.path.join(self.abs_out_path, "test.docx")
        document_docx.save(self.temp_outfile_path_docx)

        # create a docx that is really a pdf
        fake_docx = FPDF()
        fake_docx.add_page()
        fake_docx.set_xy(0, 0)
        fake_docx.set_font("arial", "B", 13.0)
        fake_docx.cell(ln=0, h=5.0, align="L", w=0, txt=text, border=0)
        self.temp_outfile_path_fake_docx = os.path.join(
            self.abs_out_path, "fake_docx.docx"
        )
        fake_docx.output(self.temp_outfile_path_fake_docx, "F")

        # create mock doc that is really rtf
        self.temp_outfile_path_fake_doc = os.path.join(
            self.abs_out_path, "fake_doc.doc"
        )
        # write in rtf format but save as doc
        rtf_text = r"{\rtf{\fonttbl {\f0 Times New Roman;}}\f0\fs60 This is a test}"
        with open(self.temp_outfile_path_fake_doc, "w") as f:
            f.write(rtf_text)

    def tearDown(self):
        if os.path.exists(self.temp_outfile_path):
            os.remove(self.temp_outfile_path)
        if os.path.exists(self.temp_outfile_path_txt):
            os.remove(self.temp_outfile_path_txt)
        if os.path.exists(self.temp_outfile_path_pdf):
            os.remove(self.temp_outfile_path_pdf)
        if os.path.exists(self.temp_outfile_path_docx):
            os.remove(self.temp_outfile_path_docx)
        if os.path.exists(self.temp_outfile_path_fake_docx):
            os.remove(self.temp_outfile_path_fake_docx)
        # These two files are programmatically created outside the test suite
        temp_outfile_path_fake_docx_as_pdf = os.path.join(
            self.abs_out_path, "fake_docx.pdf"
        )
        if os.path.exists(temp_outfile_path_fake_docx_as_pdf):
            os.remove(temp_outfile_path_fake_docx_as_pdf)
        temp_outfile_path_fake_doc_as_rtf = os.path.join(
            self.abs_out_path, "fake_doc.rtf"
        )
        if os.path.exists(temp_outfile_path_fake_doc_as_rtf):
            os.remove(temp_outfile_path_fake_doc_as_rtf)
        if os.path.exists(self.abs_out_path):
            shutil.rmtree(self.abs_out_path)

    def test_get_doc_text_txt(self):
        result = get_doc_text(self.temp_outfile_path_txt)
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_doc_text_pdf(self):
        result = get_doc_text(self.temp_outfile_path_pdf)
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_doc_text_docx(self):
        result = get_doc_text(self.temp_outfile_path_docx)
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_doc_text_fake_docx(self):
        # see GH183
        result = get_doc_text(self.temp_outfile_path_fake_docx, rm=False)
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_doc_text_fake_doc(self):
        # see GH194
        result = get_doc_text(self.temp_outfile_path_fake_doc, rm=False)
        expected = "This is a test"
        self.assertEqual(result, expected)

    @patch('textract.process')
    def test_get_doc_text_success(self, mock_process):
        # Set up mock data
        mock_process.return_value = b'This is a test document.'

        # Call the function with mock data
        file_name = 'test.doc'
        text = get_doc_text(file_name)

        # Check the results
        self.assertEqual(text, 'This is a test document.')

    @patch('textract.process')
    def test_get_doc_text_missing_file(self, mock_process):
        # Set up mock data
        mock_process.side_effect = textract.exceptions.MissingFileError('test.doc')

        # Call the function with mock data
        file_name = 'test.doc'
        text = get_doc_text(file_name)

        # Check the results
        self.assertEqual(text, '')
        self.assertTrue(mock_process.called)

    @patch('textract.process')
    def test_get_doc_text_bad_zipfile(self, mock_process):
        # Set up mock data
        mock_process.side_effect = BadZipfile('test.docx')

        # Call the function with mock data
        file_name = 'test.docx'
        text = get_doc_text(file_name)

        # Check the results
        self.assertEqual(text, '')
        self.assertTrue(mock_process.called)

    @patch('textract.process')
    def test_get_doc_text_error(self, mock_process):
        # Set up mock data
        mock_process.side_effect = Exception('Test error')

        # Call the function with mock data
        file_name = 'test.doc'
        text = get_doc_text(file_name)

        # Check the results
        self.assertEqual(text, '')
        self.assertTrue(mock_process.called)

if __name__ == "__main__":
    unittest.main()
