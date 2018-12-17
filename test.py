import unittest
from unittest.mock import patch, Mock
import os
from datetime import datetime
from utils.fbo_nightly_scraper import NightlyFBONotices
from fixtures import nightly_file, json_str, filtered_json_str, nightly_data, updated_nightly_data
from utils.get_fbo_attachments import FboAttachments
from utils.predict import Predict
from fpdf import FPDF
from docx import Document
from bs4 import BeautifulSoup
import requests
import httpretty
from fbo import main
from utils.db.db import Notice, NoticeType, Attachment, Model, now_minus_two
from utils.db.db_utils import get_db_url, session_scope, insert_updated_nightly_file, \
                              DataAccessLayer, clear_data
from utils.db.db_utils import fetch_notice_type_id, insert_model, insert_notice_types, \
                              retrain_check, get_validation_count, get_trained_amount, \
                              get_validated_untrained_amount


def exceptionCallback(request, uri, headers):
    '''
    Create a callback body that raises an exception when opened. This simulates a bad request.
    '''
    raise requests.ConnectionError('Raising a connection error for the test. You can ignore this!')


class NightlyFBONoticesTestCase(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        notice_types= ['MOD','PRESOL','COMBINE', 'AMDCSS']
        naics = ['334111', '334118', '3343', '33451', '334516', '334614', '5112',
                 '518', '54169', '54121', '5415', '54169', '61142']
        cls.nfbo = NightlyFBONotices(date=20180506, 
                                     notice_types=notice_types, 
                                     naics=naics)
        cls.file_lines = nightly_file.nightly_file

    @classmethod
    def tearDownClass(cls):
        cls.nfbo = None
        cls.file_lines = None
        cls.empty_file_lines = None

    @httpretty.activate
    def test_download_from_ftp_error(self):
         # httpretty won't work with this method, so this is intended to simulate a bad
         # connection/timeout error 
        httpretty.register_uri(httpretty.GET, 
                               uri=NightlyFBONoticesTestCase.nfbo.ftp_url, 
                               body="doesn't matter",
                               status=200)
        result = NightlyFBONoticesTestCase.nfbo.download_from_ftp()
        self.assertIsNone(result)

    def test_download_from_ftp(self):
         # in lieu of mocking out our FTP request, we'll actually go to it
         # and assert that we're able to read/return the lines as a list
        result = NightlyFBONoticesTestCase.nfbo.download_from_ftp()
        self.assertIsInstance(result, list)

    def test__id_and_count_notice_tags(self):
        result = NightlyFBONoticesTestCase.nfbo._id_and_count_notice_tags(NightlyFBONoticesTestCase.file_lines)
        expected = {'PRESOL': 1, 'COMBINE': 1, 'ARCHIVE': 1, 
                    'AWARD': 1, 'MOD': 1, 'AMDCSS': 1, 'SRCSGT': 1, 'UNARCHIVE': 1}
        self.assertDictEqual(result, expected)

    def test__merge_dicts(self):
        notice = [{'a': '123'}, {'b': '345'}, {'c': '678'}, {'c': '9'}]
        result = NightlyFBONoticesTestCase.nfbo._merge_dicts(notice)
        expected = {'a': '123', 'b': '345', 'c': '678 9'}
        self.assertDictEqual(result, expected)

    def test_pseudo_xml_to_json(self):
        result = "".join(sorted(NightlyFBONoticesTestCase.nfbo.pseudo_xml_to_json(NightlyFBONoticesTestCase.file_lines)))
        expected = "".join(sorted(json_str.json_str))
        self.assertEqual(result, expected)

    def test_filter_json(self):
        result = "".join(sorted(NightlyFBONoticesTestCase.nfbo.filter_json(json_str.json_str)))
        expected = "".join(sorted(filtered_json_str.filtered_json_str))
        self.assertEqual(result, expected)
    

class FboAttachmentsTestCase(unittest.TestCase):
    
    def setUp(self):
        self.fake_fbo_url = 'https://www.fbo.gov/fake'
        self.fboa = FboAttachments(nightly_data = nightly_data.nightly_data)
        text = "This is a test"
        temp_outfile_path = 'temp_test_file'
        with open(temp_outfile_path, 'w') as f:
            f.write(text)
        self.temp_outfile_path = temp_outfile_path
        
        txt_text = "This is a test"
        temp_outfile_path_txt = 'temp_test_file_txt.txt'
        with open(temp_outfile_path_txt, 'w') as f:
            f.write(txt_text)
        self.temp_outfile_path_txt = temp_outfile_path_txt

        pdf = FPDF()
        pdf.add_page()
        pdf.set_xy(0, 0)
        pdf.set_font('arial', 'B', 13.0)
        pdf.cell(ln=0, h=5.0, align='L', w=0, txt="This is a test", border=0)
        pdf.output('test.pdf', 'F')
        self.temp_outfile_path_pdf = 'test.pdf'

        document_docx = Document()
        document_docx.add_heading("This is a test", 0)
        document_docx.save('test.docx')
        self.temp_outfile_path_docx = 'test.docx'
        self.temp_outfile_path_doc = 'fixtures/test.doc'
        

    def tearDown(self):
        self.fboa = None
        self.fake_fbo_url = None
        if os.path.exists(self.temp_outfile_path):  
            os.remove(self.temp_outfile_path)
        if os.path.exists(self.temp_outfile_path_txt):  
            os.remove(self.temp_outfile_path_txt)
        if os.path.exists(self.temp_outfile_path_pdf):
            os.remove(self.temp_outfile_path_pdf)
        if os.path.exists(self.temp_outfile_path_docx):
            os.remove(self.temp_outfile_path_docx)

        
    @httpretty.activate
    def test_get_divs(self):
        body_with_div = b'''
                            <div class="notice_attachment_ro notice_attachment_last">
                            <div>	
                        <div class="file"><input type="hidden" name="dnf_class_values[procurement_notice_archive][packages][5][files][0][file][0][preview]" value="&lt;a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'&gt;FA852618Q0033_______0001.pdf&lt;/a&gt;"><a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'>FA852618Q0033_______0001.pdf</a> (16.54 Kb)</div>
                        </div>
                            <div><span class="label">Description:</span> Amendment 0001</div>	</div>


                            </div><!-- widget -->
                        
                            </div>
        '''
        httpretty.register_uri(httpretty.GET, 
                               uri=self.fake_fbo_url, 
                               body=body_with_div, 
                               content_type = "text/html")
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = ['expected div in here']
        self.assertEqual(len(result), len(expected))
    
    @httpretty.activate
    def test_get_divs_wrong_url(self):
        httpretty.register_uri(httpretty.GET, 
                               uri=self.fake_fbo_url, 
                               body='No divs in here')
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))
    
    @httpretty.activate
    def test_get_divs_non200_url(self):
        httpretty.register_uri(httpretty.GET, 
                               uri=self.fake_fbo_url, 
                               status=404)
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))
        
    @httpretty.activate
    def test_get_divs_connection_error(self):
        httpretty.register_uri(method=httpretty.GET, 
                               uri=self.fake_fbo_url, 
                               status=200, 
                               body=exceptionCallback)
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))
            
    @httpretty.activate
    def test_get_neco_navy_mil_attachment_urls_singleton(self):
        body = b'''
                <table id="tbl7" border="0" style="width:600px;">
                    <tbody><tr id="dwnld2_row">
                        <td class="tbl_hdr" align="right" style="width:150px;">Download File: </td><td class="tbl_itm_sm" align="left">&nbsp;<a id="dwnld2_hypr" href="/upload/N00189/N0018919Q0353Combined_Synopsis_Solicitation.docx" target="_blank">N00189/N0018919Q0353Combined_Synopsis_Solicitation.docx</a></td>
                    </tr>
                </tbody></table>
                '''
        httpretty.register_uri(httpretty.GET, 
                               uri=self.fake_fbo_url, 
                               status=200,
                               body=body, 
                               content_type = "text/html")
        result = self.fboa.get_neco_navy_mil_attachment_urls(self.fake_fbo_url)
        expected = ["https://www.neco.navy.mil/upload/N00189/N0018919Q0353Combined_Synopsis_Solicitation.docx"]
        self.assertListEqual(result, expected)
    
    @httpretty.activate
    def test_get_neco_navy_mil_attachment_urls_multiple(self):
        body = '''
                <tbody><tr id="dwnld2_row">
                        <td class="tbl_hdr" align="right" style="width:150px;">Download File: </td><td class="tbl_itm_sm" align="left">&nbsp;<a id="dwnld2_hypr" href="/upload/N00406/N0040619Q0062N00406-19-Q-0062_SOLICITATION.pdf" target="_blank">N00406/N0040619Q0062N00406-19-Q-0062_SOLICITATION.pdf</a></td>
                    </tr><tr id="dwnld3_row">
                        <td class="tbl_hdr" align="right" style="width:150px;">Download File: </td><td class="tbl_itm_sm" align="left">&nbsp;<a id="dwnld3_hypr" href="/upload/N00406/N0040619Q0062252.211-7003_ITEM_UNIQUE_IDENTIFICATION__VALUATION.docx" target="_blank">N00406/N0040619Q0062252.211-7003_ITEM_UNIQUE_IDENTIFICATION__VALUATION.docx</a></td>
                    </tr><tr id="dwnld4_row">
                        <td class="tbl_hdr" align="right" style="width:150px;">Download File: </td><td class="tbl_itm_sm" align="left">&nbsp;<a id="dwnld4_hypr" href="/upload/N00406/N0040619Q0062NAVSUP_FACTS-SP_Shipping_information.docx" target="_blank">N00406/N0040619Q0062NAVSUP_FACTS-SP_Shipping_information.docx</a></td>
                    </tr>
                </tbody>
               '''
        httpretty.register_uri(httpretty.GET, 
                               uri=self.fake_fbo_url, 
                               status=200,
                               body=body, 
                               content_type = "text/html")
        result = self.fboa.get_neco_navy_mil_attachment_urls(self.fake_fbo_url)
        expected = ['https://www.neco.navy.mil/upload/N00406/N0040619Q0062N00406-19-Q-0062_SOLICITATION.pdf',
                    'https://www.neco.navy.mil/upload/N00406/N0040619Q0062252.211-7003_ITEM_UNIQUE_IDENTIFICATION__VALUATION.docx',
                    'https://www.neco.navy.mil/upload/N00406/N0040619Q0062NAVSUP_FACTS-SP_Shipping_information.docx']
        self.assertListEqual(result, expected)

    @httpretty.activate
    def test_get_neco_navy_mil_attachment_urls_connection_error(self):
        httpretty.register_uri(method=httpretty.GET, 
                               uri=self.fake_fbo_url, 
                               status=200, 
                               body=exceptionCallback)
        result = self.fboa.get_neco_navy_mil_attachment_urls(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))

    @httpretty.activate
    def test_get_neco_navy_mil_attachment_urls_non200_url(self):
        httpretty.register_uri(httpretty.GET, 
                               uri=self.fake_fbo_url, 
                               status=404)
        result = self.fboa.get_neco_navy_mil_attachment_urls(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))

    def test_insert_attachments(self):
        text = "This is a test"
        notice = {'a': '1', 'b': '2'}
        with open(self.temp_outfile_path, 'w') as f:
            f.write(text)
        file_list = [(self.temp_outfile_path,
                      self.fake_fbo_url)]
        result = self.fboa.insert_attachments(file_list, notice)
        expected = {'a': '1',
                    'b': '2',
                    'attachments':[{'text':text, 
                                    'url':self.fake_fbo_url,
                                    'prediction':None, 
                                    'decision_boundary':None,
                                    'validation':None,
                                    'trained':False}]
        }
        self.assertDictEqual(result, expected)

    @httpretty.activate
    def test_size_check(self):
        httpretty.register_uri(httpretty.HEAD, 
                               uri=self.fake_fbo_url, 
                               body='This is less than 500MB')
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = True
        self.assertEqual(result, expected)

    @httpretty.activate
    def test_size_check_non200_url(self):
        httpretty.register_uri(httpretty.HEAD, 
                               uri=self.fake_fbo_url, 
                               status=404)
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = False
        self.assertEqual(result, expected)

    @httpretty.activate
    def test_size_check_connection_error(self):
        httpretty.register_uri(method=httpretty.HEAD, 
                               uri=self.fake_fbo_url, 
                               status=200, 
                               body=exceptionCallback)
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = False
        self.assertEqual(result, expected)
    
    def test_get_filename_from_cd(self):
        content_disposition = 'attachment; filename="test.doc"'
        result = self.fboa.get_filename_from_cd(content_disposition)
        expected = 'test.doc'
        self.assertEqual(result, expected)

    def test_get_filename_from_bad_cd(self):
        content_disposition = 'no file name in here!'
        result = self.fboa.get_filename_from_cd(content_disposition)
        expected = None
        self.assertEqual(result, expected)
    
    def test_get_file_name(self):
        attachment_url = 'https://test.pdf'
        content_type = 'application/pdf'
        result = self.fboa.get_file_name(attachment_url, content_type)
        expected = 'test.pdf'
        self.assertEqual(result, expected)
    
    def test_get_attachment_text_txt(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_txt, 'url')
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_attachment_text_pdf(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_pdf, 'url')
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_attachment_text_docx(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_docx, 'url')
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_attachment_text_doc(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_doc, 'url')
        expected = "This is a test"
        self.assertEqual(result, expected)
    
    def test_get_attachment_url_from_div(self):
        div = '<a href="/utils/view?id=798e26de983ca76f9075de687047445a"\
               target="_blank" title="Download/View FD2060-17-33119_FORM_158_00.pdf"\
               class="file">FD2060-17-33119_FORM_158_00.pdf</a>'
        div = BeautifulSoup(div, "html.parser")
        result, is_neco_navy_mil = self.fboa.get_attachment_url_from_div(div)
        expected = ['https://www.fbo.gov/utils/view?id=798e26de983ca76f9075de687047445a']
        with self.subTest():
            self.assertListEqual(result, expected)
        with self.subTest():
            self.assertFalse(is_neco_navy_mil)
    
    def test_get_attachment_url_from_div_space(self):
        div = '<a href="http://  https://www.thisisalinktoanattachment.docx"\
               target="_blank" title="Download/View FD2060-17-33119_FORM_158_00.pdf"\
               class="file">FD2060-17-33119_FORM_158_00.pdf</a>'
        div = BeautifulSoup(div, "html.parser")
        result, is_neco_navy_mil = self.fboa.get_attachment_url_from_div(div)
        expected = ['https://www.thisisalinktoanattachment.docx']
        with self.subTest():
            self.assertListEqual(result, expected)
        with self.subTest():
            self.assertFalse(is_neco_navy_mil)
    
    def test_write_attachments(self):
        #TODO don't rely on the url in the div below continuing to exist
        body_with_div = b'''
                            <div class="notice_attachment_ro notice_attachment_last">
                            <div>	
                        <div class="file"><input type="hidden" name="dnf_class_values[procurement_notice_archive][packages][5][files][0][file][0][preview]" value="&lt;a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'&gt;FA852618Q0033_______0001.pdf&lt;/a&gt;"><a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'>FA852618Q0033_______0001.pdf</a> (16.54 Kb)</div>
                        </div>
                            <div><span class="label">Description:</span> Amendment 0001</div>	</div>


                            </div><!-- widget -->
                        
                            </div>
        '''
        soup = BeautifulSoup(body_with_div, "html.parser")
        attachment_divs = soup.find_all('div', {"class": "notice_attachment_ro"})
        result = self.fboa.write_attachments(attachment_divs)
        expected = [('attachments/FA852618Q0033_______0001.pdf', 
                     'https://www.fbo.gov/utils/view?id=d95550fb782f53357ed65db571ef9186')]
        for file_url_tup in result:
            file, _ = file_url_tup
            os.remove(file)
        self.assertEqual(result, expected)
        

class PredictTestCase(unittest.TestCase):

    def setUp(self):
        json_data = updated_nightly_data.updated_nightly_data
        self.predict = Predict(json_data = json_data, 
                               best_model_path='utils/binaries/best_clf_accuracy.pkl')

    def tearDown(self):
        self.predict = None
        
    def test_transform_text(self):
        test_text = "This is a testy test that's testing transform_text"
        result = self.predict.transform_text(test_text)
        expected = 'testi test test'
        self.assertEqual(result, expected)

    def test_transform_text_none(self):
        test_text = None
        result = self.predict.transform_text(test_text)
        expected = 'none'
        self.assertEqual(result, expected)

    def test_transform_text_number(self):
        test_text = 123
        result = self.predict.transform_text(test_text)
        expected = '123'
        self.assertEqual(result, expected)

    def test_insert_predictions_top_level_keys(self):
        json_data = self.predict.insert_predictions()
        result_keys = set(json_data.keys())
        expected_keys = {'COMBINE', 'PRESOL', 'AMDCSS', 'MOD'}
        self.assertEqual(result_keys, expected_keys)

    def test_insert_predictions_bottom_level_keys(self):
        json_data = self.predict.insert_predictions()
        result_keys = set(json_data['PRESOL'][0]['attachments'][0].keys())
        expected_keys = {'trained', 'decision_boundary', 'prediction', 'text', 'validation', 'url'}
        self.assertEqual(result_keys, expected_keys)

    def test_insert_predictions_value_types(self):
        json_data = self.predict.insert_predictions()
        decision_boundary = json_data['PRESOL'][0]['attachments'][0]['decision_boundary']
        with self.subTest():
            self.assertIsInstance(decision_boundary, float)
        prediction = json_data['PRESOL'][0]['attachments'][0]['prediction']
        with self.subTest():
            self.assertIsInstance(prediction, int)

    def test_insert_predictions_compliant_insert(self):
        json_data = self.predict.insert_predictions()
        notice = json_data['PRESOL'][0]
        compliant_value = notice['compliant']
        self.assertIsInstance(compliant_value, int)

class PostgresTestCase(unittest.TestCase):
    
    def setUp(self):
        conn_string = get_db_url()
        self.predicted_nightly_data = {'AMDCSS': [{'date': '0506',
                                      'year': '18',
                                      'agency': 'department of justice',
                                      'office': 'federal bureau of investigation',
                                      'location': 'procurement section',
                                      'zip': '20535',
                                      'classcod': '70',
                                      'naics': '511210',
                                      'offadd': '935 pennsylvania avenue, n.w. washington dc 20535',
                                      'subject': 'enterprise business process management software tool',
                                      'solnbr': 'rfp-e-bpm-djf-18-0800-pr-0000828',
                                      'ntype': 'combine',
                                      'contact': 'clark kent, contracting officer, phone 5555555555, email clark.kent@daily-planet.com',
                                      'desc': '  link to document',
                                      'url': 'url',
                                      'setaside': 'n/a',
                                      'popcountry': 'us',
                                      'popzip': '20535',
                                      'popaddress': '935 pennsylvania ave. n.w. washington, dc  ',
                                      'attachments': [{'text': 'test_text_0',
                                                       'url': 'test_url_0',
                                                       'prediction': 1,
                                                       'decision_boundary': 0,
                                                       'validation': None,
                                                       'trained': False},
                                                       {'text': 'test_text_1',
                                                       'url': 'test_url_1',
                                                       'prediction': 1,
                                                       'decision_boundary': 0,
                                                       'validation': None,
                                                       'trained': False},
                                                       {'text': 'test_text_2',
                                                       'url': 'test_url_2',
                                                       'prediction': 1,
                                                       'decision_boundary': 0,
                                                       'validation': None,
                                                       'trained': False},
                                                       {'text': 'test_text_3',
                                                       'url': 'test_url_3',
                                                       'prediction': 1,
                                                       'decision_boundary': 0,
                                                       'validation': None,
                                                       'trained': False},
                                                       {'text': 'test_text_4',
                                                       'url': 'test_url_4',
                                                       'prediction': 1,
                                                       'decision_boundary': 0,
                                                       'validation': None,
                                                       'trained': False},
                                                       {'text': 'test_text_5',
                                                       'url': 'test_url_5',
                                                       'prediction': 1,
                                                       'decision_boundary': 0,
                                                       'validation': None,
                                                       'trained': False},
                                                       {'text': 'test_text_6',
                                                       'url': 'test_url_6',
                                                       'prediction': 1,
                                                       'decision_boundary': 0,
                                                       'validation': None,
                                                       'trained': False}],
                                      'compliant': 0}],
                            'MOD': [],
                            'COMBINE': [{'date': '0506',
                                         'year': '18',
                                         'agency': 'defense logistics agency',
                                         'office': 'dla acquisition locations',
                                         'location': 'dla aviation - bsm',
                                         'zip': '23297',
                                         'classcod': '66',
                                         'naics': '334511',
                                         'offadd': '334511',
                                         'subject': 'subject',
                                         'solnbr': 'spe4a618t934n',
                                         'respdate': '051418',
                                         'archdate': '06132018',
                                         'contact': 'bob.dylan@aol.com',
                                         'desc': 'test123',
                                         'url': 'test_url',
                                         'setaside': 'n/a  ',
                                         'attachments': [],
                                         'compliant': 0}],
                            'PRESOL': []}
        self.predicted_nightly_data_day_two = {'AMDCSS': [{'date': '0506',
                                               'year': '17',
                                               'agency': 'defense logistics agency',
                                               'office': 'dla acquisition locations',
                                               'location': 'dla aviation - bsm',
                                               'zip': '23297',
                                               'classcod': '66',
                                               'naics': '334511',
                                               'offadd': '334511',
                                               'subject': 'subject',
                                               'solnbr': 'spe4a618t934n',
                                               'respdate': '051418',
                                               'archdate': '06132018',
                                               'contact': 'bob.dylan@aol.com',
                                               'desc': 'test123',
                                               'url': 'test_url',
                                               'setaside': 'n/a  ',
                                               'attachments': [{'text': 'test_text_0',
                                                                'url': 'test_url_0',
                                                                'prediction': 1,
                                                                'decision_boundary': 0,
                                                                'validation': None,
                                                                'trained': False}],
                                               'compliant': 0}]
                                            }
        self.dal = DataAccessLayer(conn_string = conn_string)
        self.dal.connect()
    
    def tearDown(self):
        with session_scope(self.dal) as session:
            clear_data(session)
        self.dal.drop_local_postgres_db()
        self.dal = None
        self.predicted_nightly_data = None
        self.predicted_nightly_data_day_two = None
    
    def test_insert_notice_types(self):
        with session_scope(self.dal) as session:
            insert_notice_types(session)
            notice_types= ['MOD','PRESOL','COMBINE', 'AMDCSS', 'TRAIN']
            notice_type_ids = []
            for notice_type in notice_types:
                notice_type_id = session.query(NoticeType.id).filter(NoticeType.notice_type==notice_type).first().id
                notice_type_ids.append(notice_type_id)
            notice_type_ids = set(notice_type_ids)
            result = len(notice_type_ids)
            expected = len(notice_types)
            self.assertEqual(result, expected)
        
    def test_insert_updated_nightly_file(self):
        with session_scope(self.dal) as session:
            insert_updated_nightly_file(session, self.predicted_nightly_data)
            notice_types= ['MOD','PRESOL','COMBINE', 'AMDCSS', 'TRAIN']
            result_notice_types = []
            result_notices = []
            result_predictions = []
            notice_dates = []
            for nt in notice_types:
                n_id = fetch_notice_type_id(nt, session)
                result_notice_types.append(n_id)
                n = session.query(NoticeType).get(n_id)
                notices = n.notices
                for notice in notices:
                    notice_date = notice.date
                    notice_dates.append(notice_date)
                    result_notices.append(notice)
                    notice_attachments = notice.attachments
                    for a in notice_attachments:
                        result_predictions.append(a.prediction)
            with self.subTest():
                predictions_result = len(result_predictions)
                prediction_expected = 7
                self.assertEqual(predictions_result, prediction_expected)
            with self.subTest():
                notices_result = len(result_notices)
                notices_expected = 2
                self.assertEqual(notices_result, notices_expected)
            with self.subTest():
                notice_types_result = len(result_notice_types)
                notice_types_expected = 5
                self.assertEqual(notice_types_result, notice_types_expected)
            with self.subTest():
                notice_dates_result = set([date.strftime("%Y%m%d") for date in notice_dates])
                notice_dates_expected = set([now_minus_two().strftime("%Y%m%d")])
                self.assertSetEqual(notice_dates_result,notice_dates_expected)

    def test_insert_model(self):
        with session_scope(self.dal) as session:
            insert_model(session, estimator = 'SGDClassifier', best_params = {'a':'b'})
            model = session.query(Model).filter(Model.estimator=='SGDClassifier').first()
            result = model.estimator
            expected = 'SGDClassifier'
            self.assertEqual(result, expected)

    def test_insert_updated_nightly_file_day_two(self):
        with session_scope(self.dal) as session:
            insert_updated_nightly_file(session, self.predicted_nightly_data)
            insert_updated_nightly_file(session, self.predicted_nightly_data_day_two)
            notice_number = 'SPE4A618T934N'.lower()
            notice_ids = session.query(Notice.id).filter(Notice.notice_number==notice_number).all()
            result = len(notice_ids)
            expected = 2
            self.assertEqual(result, expected)

    def test_get_validation_count(self):
        with session_scope(self.dal) as session:
            insert_updated_nightly_file(session, self.predicted_nightly_data)
            result = get_validation_count(session)
            expected = 0
            self.assertEqual(result, expected)

    def test_get_trained_amount(self):
        with session_scope(self.dal) as session:
            insert_updated_nightly_file(session, self.predicted_nightly_data)
            result = get_trained_amount(session)
            expected = 0
            self.assertEqual(result, expected)

    def test_get_validated_untrained_amount(self):
        with session_scope(self.dal) as session:
            result = get_validated_untrained_amount(session)
            expected = 0
            self.assertEqual(result, expected)

    def test_retrain_check(self):
        with session_scope(self.dal) as session:
            insert_updated_nightly_file(session, self.predicted_nightly_data)
            result = retrain_check(session)
            expected = 0
            self.assertEqual(result, expected)

class EndToEndTest(unittest.TestCase):
    def setUp(self):
        conn_string = get_db_url()
        self.dal = DataAccessLayer(conn_string)
        self.dal.connect()
        self.main = main

    def tearDown(self):
        self.dal = None
        self.main = None
    
    @patch('utils.fbo_nightly_scraper')
    def test_main(self, fbo_mock):
        nfbo = fbo_mock.NightlyFBONotices.return_value
        # use 10/28 since the 28th's file is only 325 kB
        nfbo.ftp_url = 'ftp://ftp.fbo.gov/FBOFeed20181028'
        with self.subTest():
            self.main()
            self.assertTrue(True)
        with self.subTest():
            cwd = os.getcwd()
            attachments_dir = os.path.join(cwd, 'attachments')
            dir_exists = os.path.isdir(attachments_dir)
            self.assertFalse(dir_exists)

if __name__ == '__main__':
    unittest.main()